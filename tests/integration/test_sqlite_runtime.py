from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace

import openpyxl

from app.db.base import Base
from app.db.engine import create_database_engine, create_session_factory
from app.db.models import BatchRun, Document, DocumentReference
from app.services.analytics_service import get_dashboard_summary, get_quality_report
from app.services.export_service import export_excel


def test_sqlite_runtime_engine_pragmas_and_reporting_smoke(tmp_path):
    database_path = tmp_path / "olre.sqlite3"
    settings = SimpleNamespace(resolved_database_url=f"sqlite:///{database_path.as_posix()}")
    engine = create_database_engine(settings)

    try:
        with engine.connect() as connection:
            assert connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one() == 1
            assert connection.exec_driver_sql("PRAGMA busy_timeout").scalar_one() == 5000
            journal_mode = connection.exec_driver_sql("PRAGMA journal_mode").scalar_one()
            assert journal_mode in {"wal", "memory"}

        Base.metadata.create_all(engine)
        session_factory = create_session_factory(engine)
        with session_factory() as session:
            batch = BatchRun(triggered_by="sqlite-test", status="completed")
            session.add(batch)
            session.flush()
            document = Document(
                batch_run_id=batch.id,
                original_file_name="sqlite-smoke.pdf",
                content_hash="sqlite-smoke-hash",
                file_size_bytes=123,
                page_count=1,
                processing_status="processed",
                processed_at=datetime.now(UTC),
                moved_to_path=str(tmp_path / "sqlite-smoke.pdf"),
            )
            session.add(document)
            session.flush()
            session.add(
                DocumentReference(
                    document_id=document.id,
                    page_number=1,
                    source_type="text",
                    reference_class="url",
                    raw_reference="https://example.com/sqlite",
                    final_url="https://example.com/sqlite",
                    resolution_status="resolved",
                    http_status=200,
                )
            )
            session.commit()

        with session_factory() as session:
            dashboard = get_dashboard_summary(session)
            quality = get_quality_report(session)
            excel_response = export_excel(session, {})

        assert dashboard["total_documents"] == 1
        assert dashboard["total_references"] == 1
        assert quality["failed_documents"] == []
        workbook = openpyxl.load_workbook(BytesIO(excel_response.body))
        assert workbook.sheetnames == ["Summary", "Documents", "References", "Domains", "Errors"]
    finally:
        engine.dispose()
